"""Printable worksheets with native editable OMML equations via Pandoc."""
import copy
import json
import re
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from .config import settings
from . import exercise_store as store
from .ffmpeg_tools import VideoProcessingError

TYPE_NAMES = {"single_choice": "单选题", "fill_blank": "填空题", "short_answer": "简答题", "calculation": "计算题"}


def _pandoc(args, content):
    try:
        result = subprocess.run(["pandoc", *args], input=content, capture_output=True, timeout=60, check=False)
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise VideoProcessingError(f"Word 转换失败：{exc}") from exc
    if result.returncode or b"[WARNING]" in result.stderr:
        raise VideoProcessingError("公式或内容无法转换为 Word：" + result.stderr.decode("utf-8", errors="replace")[:800])
    return result.stdout


def _safe_ast(node):
    if isinstance(node, list):
        return [_safe_ast(item) for item in node]
    if not isinstance(node, dict):
        return node
    if node.get("t") in {"Image", "RawBlock", "RawInline"}:
        raise VideoProcessingError("题目含有不支持的图片或原始标记，请重新生成。")
    if node.get("t") == "Link":
        # Only text and mathematics are exported; no external resources/relations.
        return {"t": "Span", "c": [["", [], []], _safe_ast(node["c"][1])]}
    return {key: _safe_ast(value) for key, value in node.items()}


def _content_document(markdown):
    from docx import Document
    from io import BytesIO

    ast = json.loads(_pandoc(["-f", "markdown+tex_math_dollars-raw_html-raw_tex", "-t", "json"], markdown.encode()))
    ast = _safe_ast(ast)
    data = _pandoc(["-f", "json", "-t", "docx"], json.dumps(ast).encode())
    return Document(BytesIO(data))


def _append_content(doc, markdown, number):
    try:
        converted = _content_document(markdown)
    except VideoProcessingError as exc:
        raise VideoProcessingError(f"第 {number} 题：{exc}") from exc
    # Numbering IDs are local to each DOCX, so copy and remap definitions.
    from docx.oxml.ns import qn
    target = doc.part.numbering_part.element
    source = converted.part.numbering_part.element
    abstracts, nums = {}, {}
    next_abstract = max([int(e.get(qn("w:abstractNumId"))) for e in target if e.tag == qn("w:abstractNum")] + [-1]) + 1
    next_num = max([int(e.get(qn("w:numId"))) for e in target if e.tag == qn("w:num")] + [0]) + 1
    for element in source:
        if element.tag == qn("w:abstractNum"):
            copied = copy.deepcopy(element)
            old_id = element.get(qn("w:abstractNumId"))
            abstracts[old_id] = str(next_abstract)
            copied.set(qn("w:abstractNumId"), str(next_abstract))
            next_abstract += 1
            target.append(copied)
    for element in source:
        if element.tag == qn("w:num"):
            copied = copy.deepcopy(element)
            old_id = element.get(qn("w:numId"))
            nums[old_id] = str(next_num)
            copied.set(qn("w:numId"), str(next_num))
            next_num += 1
            reference = copied.find(qn("w:abstractNumId"))
            reference.set(qn("w:val"), abstracts[reference.get(qn("w:val"))])
            target.append(copied)
    for style in converted.styles:
        if style.style_id not in {s.style_id for s in doc.styles}:
            doc.styles.element.append(copy.deepcopy(style.element))
    for element in converted.element.body:
        if element.tag == qn("w:sectPr"):
            continue
        copied = copy.deepcopy(element)
        for reference in copied.xpath('.//*[local-name()="numPr"]/*[local-name()="numId"]'):
            reference.set(qn("w:val"), nums[reference.get(qn("w:val"))])
        doc.element.body.insert(len(doc.element.body) - 1, copied)


def export(lesson_id, request):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if not shutil.which("pandoc"):
        raise VideoProcessingError("未安装 Pandoc，无法导出 Word；请安装部署说明中的依赖。")
    lesson = store.lesson(lesson_id)
    batch = next((b for b in lesson["batches"] if b["id"] == request.batch_id), None)
    if batch is None:
        raise LookupError("题目批次不存在")
    if len(set(request.question_ids)) != len(request.question_ids):
        raise VideoProcessingError("不能重复选择题目")
    if not set(request.question_ids).issubset({q["id"] for q in batch["questions"]}):
        raise VideoProcessingError("所选题目不属于当前批次")
    # Stable batch order is identical regardless of checkbox click order.
    selected = [q for q in batch["questions"] if q["id"] in request.question_ids]
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    for style_name in ("Normal", "Body Text", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Noto Serif CJK SC"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Serif CJK SC")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12 if style_name in ("Normal", "Body Text") else 14)
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(8)
    document.styles["Title"].font.size = Pt(20)
    document.add_paragraph(request.title, "Title")
    answers = request.kind == "answers"
    document.add_paragraph("答案解析卷" if answers else "练习卷")
    if not answers:
        document.add_paragraph("姓名：________________    日期：________________")
    else:
        document.add_paragraph("题号与练习卷一致。")
    # Preserve the knowledge snapshot attached to this batch after later edits.
    points = {p["id"]: p for p in batch.get("knowledge_points", [])}
    for number, question in enumerate(selected, 1):
        document.add_paragraph(f"第 {number} 题  {TYPE_NAMES[question['type']]}", "Heading 2")
        _append_content(document, question["stem"], number)
        for index, option in enumerate(question["options"]):
            _append_content(document, f"{chr(65 + index)}、{option}", number)
        if answers:
            document.add_paragraph("参考答案", "Heading 3")
            _append_content(document, question["answer"], number)
            document.add_paragraph("解析", "Heading 3")
            _append_content(document, question["explanation"], number)
            titles = [points[i]["title"] for i in question["knowledge_point_ids"] if i in points]
            if titles:
                document.add_paragraph("对应知识点：" + "、".join(titles))
        else:
            blank = document.add_paragraph("答：" if question["type"] in ("short_answer", "calculation") else "")
            blank.paragraph_format.space_after = Pt({"single_choice": 12, "fill_blank": 24, "short_answer": 100, "calculation": 150}[question["type"]])
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.widow_control = True
        # Let lengthy worked solutions span pages without losing whole pages.
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")
    filename = re.sub(r"[^\w\-\u4e00-\u9fff]", "_", request.title)[:60]
    filename += ("_答案解析_" if answers else "_练习卷_") + uuid.uuid4().hex[:10] + ".docx"
    settings.outputs_dir.mkdir(parents=True, exist_ok=True)
    output = settings.outputs_dir / filename
    # Atomic save avoids leaving a half-written downloadable file on failure.
    with tempfile.TemporaryDirectory(dir=settings.outputs_dir) as temp:
        intermediate = Path(temp) / "document.docx"
        document.save(intermediate)
        intermediate.replace(output)
    return {"filename": filename, "download_url": f"/api/outputs/{filename}", "size_bytes": output.stat().st_size}
