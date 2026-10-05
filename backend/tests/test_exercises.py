"""Exercise pipeline, persistence, concurrency, and native Word export checks."""
import json
from pathlib import Path
import threading
import time
from types import SimpleNamespace
from unittest.mock import patch
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from pydantic import ValidationError

from backend.app.config import settings
from backend.app import exercise_service as service, exercise_store as store, exercise_export as exporter, storage
from backend.app import asr_service
from backend.app.exercise_models import GenerateRequest, LessonPatch, Question, ExportExercisesRequest, SaveExercisePageRequest
from backend.app.exercise_routes import router
from backend.app.ffmpeg_tools import VideoProcessingError


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    for name in ("exercises_dir", "uploads_dir", "downloads_dir", "outputs_dir", "whisper_models_dir"):
        directory = tmp_path / name
        directory.mkdir()
        monkeypatch.setattr(settings, name, directory)
    monkeypatch.setattr(settings, "openai_api_key", "test-key")
    monkeypatch.setattr(settings, "asr_backend", "whisper")
    monkeypatch.setattr(asr_service, "_model", None)
    monkeypatch.setattr(asr_service, "_model_key", None)
    service.start()
    yield
    service.stop()


def lesson(segments=True):
    value = {"id": "lesson", "video_id": "a" * 32, "title": "代数课", "duration": 100, "source": {"type": "upload", "bv": None, "page": None}, "version": 1, "segments": [{"start": 0, "end": 10, "text": "二次方程的求根公式"}] if segments else [], "knowledge_points": [{"id": "point", "title": "二次方程", "detail": "求根公式与判别式", "formulas": "$x^2$", "segment_ids": [0]}] if segments else [], "batches": [], "latest_job_id": None}
    store.save_lesson(value)
    return value


def question(**updates):
    value = {"id": "q", "type": "calculation", "difficulty": "practice", "stem": "求解 $x^2=4$。", "options": [], "answer": "$x=\\pm 2$", "explanation": "两边开平方，得到正负两个根。", "knowledge_point_ids": ["point"]}
    value.update(updates)
    return value


def wait_job(task):
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        job = store.job(task["job_id"])
        if job["status"] in ("succeeded", "failed"):
            return job
        time.sleep(.01)
    raise AssertionError("job did not finish")


def test_restart_marks_only_unfinished_jobs():
    lesson()
    for status in ("queued", "running", "succeeded"):
        store.save_job({"id": status, "lesson_id": "lesson", "status": status})
    store.init_db()
    assert store.job("queued")["status"] == "failed"
    assert store.job("running")["stage"] == "interrupted"
    assert store.job("succeeded")["status"] == "succeeded"


def test_long_segments_chunk_without_losing_tail_or_evidence(monkeypatch):
    monkeypatch.setattr(settings, "exercise_max_input_chars", 1000)
    segments = [{"start": 0, "end": 10, "text": "甲" * 2500}, {"start": 10, "end": 20, "text": "最后知识点"}]
    chunks = service._source_chunks(segments)
    assert len(chunks) >= 3
    assert sum(text.count("甲") for text, ids in chunks) == 2500
    assert "最后知识点" in chunks[-1][0]
    assert 1 in chunks[-1][1]


def test_knowledge_extraction_merges_evidence_and_rejects_wrong_ids(monkeypatch):
    segments = [{"start": 0, "end": 10, "text": "甲" * 1000}, {"start": 10, "end": 20, "text": "乙" * 1000}]
    monkeypatch.setattr(settings, "exercise_max_input_chars", 1100)
    calls = iter([json.dumps({"knowledge_points": [{"id": "a", "title": "概念", "detail": "前半段", "segment_ids": [0]}]}, ensure_ascii=False), json.dumps({"knowledge_points": [{"id": "a", "title": "概念", "detail": "后半段", "segment_ids": [1]}]}, ensure_ascii=False)])
    with patch.object(service, "_chat", side_effect=lambda *args, **kwargs: next(calls)):
        points = service.extract_knowledge(segments, {"id": "j", "lesson_id": "lesson"})
    assert len(points) == 1
    assert points[0]["segment_ids"] == [0, 1]
    assert "后半段" in points[0]["detail"]


def test_structured_output_repairs_once_and_never_accepts_prose(caplog):
    with patch.object(service, "_chat", side_effect=["not json", '{"value": 42}']) as chat:
        assert service._structured("prompt", lambda p: p["value"]) == 42
        assert chat.call_count == 2
        assert "not a network retry" in caplog.text
    with patch.object(service, "_chat", return_value="broken") as chat:
        with pytest.raises(VideoProcessingError):
            service._structured("prompt", lambda p: p)
        assert chat.call_count == 2


@pytest.mark.parametrize("updates", [{"type": "single_choice", "options": ["one"], "answer": "A"}, {"type": "single_choice", "options": ["one", "two", "three", "four"], "answer": "two"}, {"options": ["unexpected"]}, {"answer": ""}])
def test_question_shape_requires_valid_answers_and_options(updates):
    with pytest.raises(ValidationError):
        Question.model_validate(question(**updates))


def test_no_duplicate_or_unselected_questions():
    request = GenerateRequest(version=1, knowledge_point_ids=["point"], types=["calculation"], count=2)
    with pytest.raises(ValueError, match="重复"):
        service._validate_questions({"questions": [question(), question()]}, 2, {"point"}, request)
    with pytest.raises(ValueError, match="未选择"):
        service._validate_questions({"questions": [question(knowledge_point_ids=["else"])]}, 1, {"point"}, request)


def test_generation_review_keeps_requested_count_and_snapshot():
    lesson()
    request = GenerateRequest(version=1, knowledge_point_ids=["point"], types=["calculation"], count=2)
    qs = [question(id="a", stem="题一 $x^2=4$"), question(id="b", stem="题二 $x^2=9$")]
    with patch.object(service, "_chat", return_value=json.dumps({"questions": qs})) as chat:
        task = service.generate("lesson", request)
        assert wait_job(task)["status"] == "succeeded"
    saved = store.lesson("lesson")
    assert chat.call_count == 2  # Generation plus answer/solution review.
    batch = saved["batches"][0]
    assert len(batch["questions"]) == 2
    assert batch["knowledge_points"] == saved["knowledge_points"]
    assert len({q["id"] for q in batch["questions"]}) == 2
    page = store.page(batch["page_slug"])
    assert page["slug"] == "er-ci-fang-cheng"
    assert page["questions"] == batch["questions"]
    assert store.recent_pages()["total"] == 1


@pytest.mark.parametrize("count", [10, 30])
def test_generation_merges_points_preserves_assignments_and_carries_history(count):
    value = lesson()
    value["knowledge_points"] = [dict(value["knowledge_points"][0], id=f"point-{index}") for index in range(count)]
    request = GenerateRequest(version=1, knowledge_point_ids=[p["id"] for p in value["knowledge_points"]], count=count)
    calls, responses = [], []
    def chat(prompt, max_tokens, *, history):
        calls.append({"prompt": prompt, "history": json.loads(json.dumps(history))})
        if len(calls) % 2 == 1:
            assignments = json.loads(prompt.splitlines()[2])
            assert len(assignments) == 5
            # Shared transcript evidence appears only once in this prompt.
            assert prompt.count(value["segments"][0]["text"]) == 1
            questions = []
            for assignment in assignments:
                point_id = assignment["knowledge_point_ids"][0]
                kind = assignment["type"]
                questions.append(question(
                    id=point_id, type=kind, difficulty=request.difficulty,
                    stem=f"{point_id}（________）" if kind == "fill_blank" else f"{point_id} 题干",
                    options=["一", "二", "三", "四"] if kind == "single_choice" else [],
                    answer="A" if kind == "single_choice" else f"{point_id} 的完整答案",
                    explanation=f"{point_id} 的完整解析", knowledge_point_ids=[point_id],
                ))
            response = json.dumps({"questions": questions}, ensure_ascii=False)
        else:
            assert history[-1]["role"] == "assistant"
            previous = json.loads(history[-1]["content"])
            # Distinguish the reviewed result from the initial candidates.
            for q in previous["questions"]:
                q["explanation"] += "（已复核）"
            response = json.dumps(previous, ensure_ascii=False)
        if len(calls) > 1:
            assert [m["content"] for m in history if m["role"] == "assistant"] == responses
        responses.append(response)
        return response
    with patch.object(service, "_chat", side_effect=chat):
        result = service.generate_questions(value, request, {"id": "batch-job", "lesson_id": value["id"]})
    assert len(calls) == 2 * (count // 5)
    assert len(result) == count
    assert [q["knowledge_point_ids"] for q in result] == [[p["id"]] for p in value["knowledge_points"]]
    assert [q["type"] for q in result] == [request.types[index % len(request.types)] for index in range(count)]
    assert all(q["explanation"].endswith("（已复核）") for q in result)
    assert len({q["id"] for q in result}) == count
    assert calls[0]["history"] == [{"role": "system", "content": service.RULES}]


def test_structured_repair_sees_invalid_response_and_only_remembers_success():
    history = [{"role": "system", "content": service.RULES}]
    calls = []
    def chat(prompt, max_tokens, *, history):
        calls.append((prompt, json.loads(json.dumps(history))))
        return "broken output" if len(calls) == 1 else '{"value": 42}'
    with patch.object(service, "_chat", side_effect=chat):
        assert service._structured("原始出题请求", lambda p: p["value"], history=history) == 42
    assert calls[1][1][-2:] == [
        {"role": "user", "content": "原始出题请求"},
        {"role": "assistant", "content": "broken output"},
    ]
    assert "不是有效 JSON" in calls[1][0]
    assert history == [
        {"role": "system", "content": service.RULES},
        {"role": "user", "content": "原始出题请求"},
        {"role": "assistant", "content": '{"value": 42}'},
    ]


def test_later_batch_duplicate_is_repaired_using_previous_questions():
    value = lesson()
    request = GenerateRequest(version=1, knowledge_point_ids=["point"], types=["calculation"], count=6)
    first = [question(id=f"q{index}", stem=f"第 {index} 道题") for index in range(5)]
    repeated = [question(id="repeat", stem=first[0]["stem"])]
    last = [question(id="last", stem="新的第六题")]
    responses = [first, first, repeated, last, last]
    calls = []
    def chat(prompt, max_tokens, *, history):
        calls.append((prompt, json.loads(json.dumps(history))))
        return json.dumps({"questions": responses[len(calls) - 1]}, ensure_ascii=False)
    with patch.object(service, "_chat", side_effect=chat):
        result = service.generate_questions(value, request, {"id": "duplicate-job", "lesson_id": value["id"]})
    assert len(calls) == 5
    assert "题目重复" in calls[3][0]
    assert json.loads(calls[3][1][-1]["content"])["questions"] == repeated
    assert [json.loads(m["content"])["questions"] for m in calls[4][1] if m["role"] == "assistant"] == [first, first, last]
    assert len({q["stem"] for q in result}) == 6


def test_question_batches_split_large_evidence_without_dropping_points(monkeypatch):
    value = lesson()
    monkeypatch.setattr(settings, "exercise_max_input_chars", 1000)
    value["segments"] = [{"start": index, "end": index + 1, "text": str(index) * 1500} for index in range(5)]
    value["knowledge_points"] = [dict(value["knowledge_points"][0], id=f"p{index}", segment_ids=[index]) for index in range(5)]
    request = GenerateRequest(version=1, knowledge_point_ids=[p["id"] for p in value["knowledge_points"]], count=5)
    batches = service._question_batches(value, value["knowledge_points"], request)
    assert len(batches) == 5
    assert [a["knowledge_point_ids"] for batch in batches for a in batch] == [[f"p{index}"] for index in range(5)]
    assert all(len(service._question_prompt(value, value["knowledge_points"], batch, request.difficulty)) <= 3000 for batch in batches)


@pytest.mark.parametrize("wrong_field", ["type", "knowledge_point_ids"])
def test_merged_batch_rejects_wrong_assignment_and_failed_history_stays_local(wrong_field):
    value = lesson()
    value["knowledge_points"].append(dict(value["knowledge_points"][0], id="other"))
    request = GenerateRequest(version=1, knowledge_point_ids=["point", "other"], types=["calculation", "short_answer"], count=2)
    valid = [question(), question(id="q2", type="short_answer", stem="第二题", knowledge_point_ids=["other"])]
    invalid = json.loads(json.dumps(valid))
    if wrong_field == "type":
        invalid[0]["type"] = "short_answer"
    else:
        invalid[0]["knowledge_point_ids"] = ["other"]
    with patch.object(service, "_chat", return_value=json.dumps({"questions": invalid})):
        with pytest.raises(VideoProcessingError, match="分配|顺序"):
            service.generate_questions(value, request, {"id": "failed-job", "lesson_id": value["id"]})
    with patch.object(service, "_chat", return_value=json.dumps({"questions": valid})) as chat:
        assert len(service.generate_questions(value, request, {"id": "fresh-job", "lesson_id": value["id"]})) == 2
        assert chat.call_args_list[0].kwargs["history"] == [{"role": "system", "content": service.RULES}]


def saved_batch(value, batch_id="batch", created_at=100):
    batch = {"id": batch_id, "version": value["version"], "created_at": created_at,
             "knowledge_points": json.loads(json.dumps(value["knowledge_points"])),
             "questions": [question(id=f"{batch_id}-q1"), question(id=f"{batch_id}-q2", stem="求解 $x^2=9$。") ]}
    value["batches"].append(batch)
    store.save_lesson(value)
    return batch


def test_pages_use_pinyin_collisions_and_paginated_recent_history():
    value = lesson()
    first = saved_batch(value, "first", 100)
    second = saved_batch(value, "second", 200)
    assert first["page_slug"] == "er-ci-fang-cheng"
    assert second["page_slug"] == "er-ci-fang-cheng-2"
    recent = store.recent_pages(limit=1)
    assert recent["total"] == 2
    assert recent["items"][0]["slug"] == second["page_slug"]
    assert "questions" not in recent["items"][0]
    assert store.recent_pages(limit=1, offset=1)["items"][0]["slug"] == first["page_slug"]
    assert store.recent_pages(offset=2)["items"] == []


def test_saved_pages_survive_lesson_changes_restart_and_missing_video():
    value = lesson()
    batch = saved_batch(value)
    original = store.page(batch["page_slug"])
    value["knowledge_points"][0]["title"] = "新的知识点名称"
    value["batches"][0]["questions"][0]["stem"] = "更改后的内容"
    value["version"] = 2
    store.save_lesson(value)
    service.stop()
    service.start()
    assert store.page(batch["page_slug"]) == original
    assert store.recent_pages()["total"] == 1
    assert not (settings.uploads_dir / value["video_id"]).exists()


def test_save_selected_page_is_idempotent_and_keeps_batch_order():
    value = lesson()
    batch = saved_batch(value)
    request = SaveExercisePageRequest(batch_id=batch["id"], question_ids=["batch-q2", "batch-q1"], title="数学练习")
    first = store.create_page(value["id"], request)
    assert [q["id"] for q in first["questions"]] == ["batch-q1", "batch-q2"]
    assert store.create_page(value["id"], request) == first
    selected = store.create_page(value["id"], request.model_copy(update={"question_ids": ["batch-q2"]}))
    assert selected["question_count"] == 1
    assert selected["questions"][0]["id"] == "batch-q2"
    assert selected["slug"] != first["slug"]
    assert store.recent_pages()["total"] == 3


def test_simultaneous_page_saves_allocate_stable_unique_slugs():
    value = lesson()
    saved_batch(value)
    requests = [SaveExercisePageRequest(batch_id="batch", question_ids=["batch-q1"], title=f"练习 {i}") for i in range(6)]
    with ThreadPoolExecutor(max_workers=6) as pool:
        pages = list(pool.map(lambda request: store.create_page(value["id"], request), requests))
    assert len({p["slug"] for p in pages}) == 6
    assert store.recent_pages()["total"] == 7
    assert store.create_page(value["id"], requests[0])["slug"] == pages[0]["slug"]


def test_startup_backfills_legacy_batches_once_using_their_knowledge_snapshot():
    value = lesson()
    saved_batch(value)
    value["knowledge_points"][0]["title"] = "新的课程名称"
    value["batches"][0].pop("page_slug")
    # Simulate the database before printable pages existed.
    with store.connect() as conn:
        conn.execute("DROP TABLE pages")
        conn.execute("UPDATE lessons SET payload=? WHERE id=?", (json.dumps(value), value["id"]))
    store.init_db()
    batch = store.lesson(value["id"])["batches"][0]
    assert batch["page_slug"] == "er-ci-fang-cheng"
    assert store.page(batch["page_slug"])["knowledge_points"][0]["title"] == "二次方程"
    store.init_db()
    assert store.recent_pages()["total"] == 1


def test_page_routes_validate_selection_and_static_routes_precede_lesson_ids():
    value = lesson()
    batch = saved_batch(value)
    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    assert client.get("/api/exercises/pages").json()["total"] == 1
    assert client.get(f"/api/exercises/pages/{batch['page_slug']}").json()["question_count"] == 2
    assert client.get("/api/exercises/pages/missing").status_code == 404
    assert client.get("/api/exercises/pages?limit=101").status_code == 422
    assert client.get("/api/exercises/pages?offset=-1").status_code == 422
    body = {"batch_id": "batch", "question_ids": ["batch-q1"], "title": "打印测试"}
    assert client.post("/api/exercises/lesson/pages", json=body).status_code == 200
    for ids in (["missing"], ["batch-q1", "batch-q1"]):
        assert client.post("/api/exercises/lesson/pages", json={**body, "question_ids": ids}).status_code == 400
    for updates in ({"title": " "}, {"question_ids": []}):
        assert client.post("/api/exercises/lesson/pages", json={**body, **updates}).status_code == 422
    assert client.post("/api/exercises/lesson/pages", json={**body, "batch_id": "missing"}).status_code == 404


def test_duplicate_generation_requests_reuse_active_task_and_block_edit():
    lesson()
    entered, release = threading.Event(), threading.Event()
    def generate(*args):
        entered.set()
        assert release.wait(3)
        return [question()]
    request = GenerateRequest(version=1, knowledge_point_ids=["point"], count=1)
    with patch.object(service, "generate_questions", side_effect=generate):
        task = service.generate("lesson", request)
        assert entered.wait(1)
        try:
            assert service.generate("lesson", request) == task
            with pytest.raises(HTTPException) as conflict:
                service.patch_lesson("lesson", LessonPatch(version=1, knowledge_points=store.lesson("lesson")["knowledge_points"]))
            assert conflict.value.status_code == 409
        finally:
            release.set()
        assert wait_job(task)["status"] == "succeeded"


def test_transcript_changes_invalidate_points_and_reuse_asr_result():
    value = lesson()
    value["batches"] = [{"id": "old", "version": 1, "questions": [question()]}]
    store.save_lesson(value)
    revised = [{"start": 0, "end": 10, "text": "修正后的公式"}]
    with patch.object(service, "extract_knowledge", return_value=value["knowledge_points"]), patch.object(service, "transcribe") as asr:
        task = service.patch_lesson("lesson", LessonPatch(version=1, segments=revised))
        assert wait_job(task)["status"] == "succeeded"
        asr.assert_not_called()
    saved = store.lesson("lesson")
    assert saved["version"] == 2 and saved["segments"] == revised
    assert saved["batches"][0]["version"] == 1
    with pytest.raises(HTTPException) as stale:
        service.generate("lesson", GenerateRequest(version=1, knowledge_point_ids=["point"]))
    assert stale.value.status_code == 409


def test_prepare_retries_extraction_without_source_file_or_asr():
    value = lesson()
    value["knowledge_points"] = []
    store.save_lesson(value)
    with patch.object(service, "extract_knowledge", side_effect=VideoProcessingError("LLM down")), patch.object(service, "transcribe") as asr:
        failed = service.prepare(value["video_id"])
        assert wait_job(failed)["status"] == "failed"
        asr.assert_not_called()
    with patch.object(service, "extract_knowledge", return_value=[{"id": "point", "title": "公式", "detail": "讲解", "formulas": "", "segment_ids": [0]}]):
        assert wait_job(service.prepare(value["video_id"]))["status"] == "succeeded"


def test_pin_protects_old_source_until_unpinned():
    video_id = "a" * 32
    directory = settings.uploads_dir / video_id
    directory.mkdir()
    source = directory / "source.mp4"
    source.write_bytes(b"video")
    (directory / "meta.json").write_text(json.dumps({"source_path": str(source), "last_used_at": 0}))
    storage.pin_video(video_id)
    assert storage.delete_old_videos(0, now=time.time() + 100) == 0
    storage.unpin_video(video_id)
    assert storage.delete_old_videos(0, now=time.time() + 100) == 1


def test_asr_consumes_lazy_segments_on_cpu_and_reports_empty_speech():
    model = SimpleNamespace(transcribe=lambda *a, **k: (iter([SimpleNamespace(start=0, end=2, text="你好")]), SimpleNamespace(duration=3)))
    with patch("faster_whisper.WhisperModel", return_value=model) as constructor:
        result = service.transcribe(Path("video.mp4"), {"id": "asr", "lesson_id": "lesson"})
    assert result == [{"start": 0., "end": 2., "text": "你好"}]
    assert constructor.call_args.kwargs["device"] == "cpu"
    assert constructor.call_args.kwargs["compute_type"] == "int8"
    model.transcribe = lambda *a, **k: (iter([]), SimpleNamespace(duration=3))
    with pytest.raises(VideoProcessingError, match="有效语音"):
        service.transcribe(Path("video.mp4"), {"id": "asr", "lesson_id": "lesson"})


def test_model_download_failure_reports_endpoint_and_can_retry():
    job = {"id": "asr", "lesson_id": "lesson"}
    failure = ConnectionError("Network is unreachable")
    with patch("faster_whisper.WhisperModel", side_effect=failure):
        with pytest.raises(VideoProcessingError, match="HF_ENDPOINT") as error:
            service.transcribe(Path("video.mp4"), job)
    assert settings.hf_endpoint in str(error.value)
    assert "WHISPER_MODEL" in str(error.value)
    assert error.value.__cause__ is failure
    assert asr_service._model is None
    model = SimpleNamespace(transcribe=lambda *a, **k: (
        iter([SimpleNamespace(start=0, end=2, text="重试成功")]), SimpleNamespace(duration=3)
    ))
    with patch("faster_whisper.WhisperModel", return_value=model) as constructor:
        assert service.transcribe(Path("video.mp4"), job)[0]["text"] == "重试成功"
    assert constructor.call_args.kwargs["download_root"] == str(settings.whisper_models_dir)


def test_http_contract_and_validation():
    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    lesson()
    assert client.get("/api/exercises/lesson").status_code == 200
    assert client.get("/api/exercises/missing").status_code == 404
    assert client.post("/api/exercises/lesson/generate", json={"version": 1, "knowledge_point_ids": ["missing"]}).status_code == 400
    assert client.post("/api/exercises/lesson/generate", json={"version": 1, "knowledge_point_ids": ["point"], "count": 31}).status_code == 422
    with patch.object(settings, "openai_api_key", ""):
        assert client.post("/api/exercises/prepare", json={"video_id": "a" * 32}).status_code == 503


def require_pandoc(monkeypatch):
    import shutil
    if shutil.which("pandoc"):
        return
    try:
        import pypandoc
        import os
        monkeypatch.setenv("PATH", str(Path(pypandoc.get_pandoc_path()).parent) + os.pathsep + os.environ["PATH"])
    except ImportError:
        pytest.fail("Install Pandoc (or pypandoc_binary) to run the Word integration checks")


def test_two_word_documents_have_matching_numbers_and_native_math(monkeypatch):
    require_pandoc(monkeypatch)
    from docx import Document
    value = lesson()
    qs = [question(id="q1"), question(id="q2", stem=r"计算 $\frac{1}{2}+\sqrt{4}$ 和 $\sum_{i=1}^{3} i$。", explanation="1. 先计算根号\n2. 再计算求和", answer=r"$\frac{17}{2}$"), question(id="q3", stem="不应导出的题目")]
    value["batches"] = [{"id": "batch", "version": 1, "knowledge_points": value["knowledge_points"], "questions": qs}]
    store.save_lesson(value)
    paths = []
    for kind in ("worksheet", "answers"):
        response = exporter.export("lesson", ExportExercisesRequest(batch_id="batch", question_ids=["q2", "q1"], title="代数练习", kind=kind))
        paths.append(settings.outputs_dir / response["filename"])
    worksheet, answers = [Document(p) for p in paths]
    for document in (worksheet, answers):
        text = "\n".join(p.text for p in document.paragraphs)
        assert "第 1 题" in text and "第 2 题" in text and "第 3 题" not in text
        assert "不应导出的题目" not in text
        assert "求解" in text and "计算" in text
        assert document.sections[0].page_width.cm == pytest.approx(21, abs=.01)
    assert "参考答案" not in "\n".join(p.text for p in worksheet.paragraphs)
    assert "参考答案" in "\n".join(p.text for p in answers.paragraphs)
    for path in paths:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode()
            assert "<m:oMath" in xml and "<m:f>" in xml and "<m:rad>" in xml
            assert "Noto Serif CJK SC" in archive.read("word/styles.xml").decode()


@pytest.mark.parametrize("formula", [r"$x_i^2$", r"$\int_0^1 x\,dx$", r"$$\begin{pmatrix}1&2\\3&4\end{pmatrix}$$", r"$$f(x)=\begin{cases}x&x>0\\0&x\leq0\end{cases}$$"])
def test_common_latex_exports_as_omml(monkeypatch, formula):
    require_pandoc(monkeypatch)
    assert exporter._content_document(formula).element.body.xpath('.//*[local-name()="oMath"]')


def test_export_rejects_unsupported_formula_and_external_images(monkeypatch):
    require_pandoc(monkeypatch)
    with pytest.raises(VideoProcessingError):
        exporter._content_document(r"$\notarealcommand{x}$")
    with pytest.raises(VideoProcessingError, match="图片"):
        exporter._content_document("![x](https://example.com/private.png)")


def test_new_upload_prepare_is_cached_and_releases_source_pin():
    video_id = "b" * 32
    directory = settings.uploads_dir / video_id
    directory.mkdir()
    source = directory / "source.mp4"
    source.write_bytes(b"fake video")
    (directory / "meta.json").write_text(json.dumps({"source_path": str(source), "filename": "lesson.mp4", "duration": 100, "source_type": "upload"}))
    entered, release = threading.Event(), threading.Event()
    def transcribe(*args):
        entered.set()
        assert release.wait(3)
        return [{"start": 0, "end": 10, "text": "课程原文"}]
    points = [{"id": "point", "title": "概念", "detail": "讲解", "formulas": "", "segment_ids": [0]}]
    with patch.object(service, "transcribe", side_effect=transcribe) as asr, patch.object(service, "extract_knowledge", return_value=points):
        task = service.prepare(video_id)
        assert entered.wait(1)
        try:
            assert service.prepare(video_id) == task
            assert video_id in storage._source_pins
        finally:
            release.set()
        assert wait_job(task)["status"] == "succeeded"
        assert service.prepare(video_id) == task
        assert asr.call_count == 1
    service.stop()
    assert video_id not in storage._source_pins
    assert store.lesson(task["lesson_id"])["segments"][0]["text"] == "课程原文"


def test_fill_blank_question_requires_blank_position():
    with pytest.raises(ValidationError, match="填空位置"):
        Question.model_validate(question(type="fill_blank", stem="解释这个定义"))
    assert Question.model_validate(question(type="fill_blank", stem="公式是（________）。"))


def test_asr_options_and_client_cannot_override_server_selection(monkeypatch):
    monkeypatch.setattr(settings, "asr_backend", "sensevoice")
    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    options = client.get("/api/exercises/asr/options").json()
    assert options["default_backend"] == "sensevoice"
    assert {option["id"] for option in options["backends"]} == {"sensevoice", "qwen3", "whisper"}
    for backend in ("unknown", "qwen3"):
        response = client.post("/api/exercises/prepare", json={"video_id": "a" * 32, "asr_backend": backend})
        assert response.status_code == 422


def test_changed_server_model_is_used_for_saved_video(monkeypatch):
    value = lesson()
    monkeypatch.setattr(settings, "asr_backend", "qwen3")
    with patch.object(service, "pin_video", return_value=Path("source.mp4")), patch.object(service, "unpin_video"), patch.object(service, "transcribe", return_value=value["segments"]) as transcribe, patch.object(service, "extract_knowledge", return_value=value["knowledge_points"]):
        task = service.prepare(value["video_id"])
        assert wait_job(task)["status"] == "succeeded"
        assert transcribe.call_args.args[1]["request"]["asr_backend"] == "qwen3"
        assert service.prepare(value["video_id"]) == task
        assert transcribe.call_count == 1
    assert store.lesson(value["id"])["asr_backend"] == "qwen3"


def test_switch_model_invalidates_transcript_but_preserves_old_batches():
    value = lesson()
    value["batches"] = [{"id": "old", "version": 1, "questions": [question()]}]
    store.save_lesson(value)
    new_segments = [{"start": 2, "end": 10, "text": "千问识别结果"}]
    with patch.object(service, "pin_video", return_value=Path("source.mp4")), patch.object(service, "unpin_video") as unpin, patch.object(service, "transcribe", return_value=new_segments) as transcribe, patch.object(service, "extract_knowledge", return_value=value["knowledge_points"]):
        task = service.prepare(value["video_id"], "qwen3")
        assert wait_job(task)["status"] == "succeeded"
        assert service.prepare(value["video_id"], "qwen3") == task
        assert transcribe.call_count == 1
        assert transcribe.call_args.args[1]["request"]["asr_backend"] == "qwen3"
    unpin.assert_called_once_with(value["video_id"])
    saved = store.lesson(value["id"])
    assert saved["segments"] == new_segments
    assert saved["asr_backend"] == "qwen3"
    assert saved["version"] == 2
    assert saved["batches"][0]["version"] == 1


def test_failed_model_switch_preserves_old_content_and_can_return_old_cache():
    value = lesson()
    with patch.object(service, "pin_video", return_value=Path("source.mp4")), patch.object(service, "unpin_video"), patch.object(service, "transcribe", side_effect=VideoProcessingError("model load failed")):
        assert wait_job(service.prepare(value["video_id"], "sensevoice"))["status"] == "failed"
    saved = store.lesson(value["id"])
    assert saved["segments"] == value["segments"]
    assert saved["knowledge_points"] == value["knowledge_points"]
    assert saved["version"] == 1
    assert service.prepare(value["video_id"], "whisper") == {"lesson_id": value["id"], "job_id": None}


def test_model_switch_with_expired_source_keeps_old_content():
    value = lesson()
    with pytest.raises(VideoProcessingError):
        service.prepare(value["video_id"], "qwen3")
    assert store.lesson(value["id"]) == value


def test_retry_after_switch_reuses_new_transcript_if_extraction_failed():
    value = lesson()
    new_segments = [{"start": 2, "end": 10, "text": "SenseVoice识别结果"}]
    with patch.object(service, "pin_video", return_value=Path("source.mp4")), patch.object(service, "unpin_video"), patch.object(service, "transcribe", return_value=new_segments) as transcribe, patch.object(service, "extract_knowledge", side_effect=VideoProcessingError("LLM error")):
        assert wait_job(service.prepare(value["video_id"], "sensevoice"))["status"] == "failed"
        assert transcribe.call_count == 1
    saved = store.lesson(value["id"])
    assert saved["asr_backend"] == "sensevoice" and saved["segments"] == new_segments
    assert not saved["knowledge_points"]
    with patch.object(service, "pin_video") as pin, patch.object(service, "transcribe") as transcribe, patch.object(service, "extract_knowledge", return_value=value["knowledge_points"]):
        assert wait_job(service.prepare(value["video_id"], "sensevoice"))["status"] == "succeeded"
        pin.assert_not_called()
        transcribe.assert_not_called()
    assert store.lesson(value["id"])["version"] == 2


def test_running_task_keeps_its_model_and_rejects_switch(monkeypatch):
    value = lesson(segments=False)
    entered, release = threading.Event(), threading.Event()
    def transcribe(path, job):
        entered.set()
        assert release.wait(3)
        assert job["request"]["asr_backend"] == "sensevoice"
        return [{"start": 0, "end": 1, "text": "原文"}]
    with patch.object(service, "pin_video", return_value=Path("source.mp4")), patch.object(service, "unpin_video"), patch.object(service, "transcribe", side_effect=transcribe), patch.object(service, "extract_knowledge", return_value=[{"id": "p", "title": "概念", "detail": "讲解", "formulas": "", "segment_ids": [0]}]):
        task = service.prepare(value["video_id"], "sensevoice")
        assert entered.wait(1)
        try:
            monkeypatch.setattr(settings, "asr_backend", "qwen3")
            assert service.prepare(value["video_id"], "sensevoice") == task
            with pytest.raises(HTTPException) as conflict:
                service.prepare(value["video_id"], "qwen3")
            assert conflict.value.status_code == 409
        finally:
            release.set()
        assert wait_job(task)["status"] == "succeeded"
